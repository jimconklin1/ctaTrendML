import numpy as np
import pandas as pd
import itertools
import argparse
import os
import operator

import openpyxl.chart as opxc
import openpyxl.utils.datetime as opxd
import openpyxl.chart.axis as opxca
import openpyxl.drawing.line as opxdl
import openpyxl.chart.shapes as opxcs
import openpyxl.styles.alignment as opxsa
import openpyxl.styles as opxs
import openpyxl.utils as opxu

from openpyxl.utils.dataframe import dataframe_to_rows

from .basic_util import *
from .opx_util import *

SMALL_FLOAT = 1e-8

def empty_pos():
    res = argparse.Namespace(
        stat = argparse.Namespace(hold_lengths=[], exit_delays=[])
    )
    return res

pad_pos_array = lambda x, n=1: np.pad(x, ((n,0),(0,0)))
shift_pos_array = lambda x, n=1: pad_pos_array(x, n)[:-n, :]

def pos_shift(x, n=1):
    res = empty_pos()
    res.pos = shift_pos_array(x.pos, n)
    res.seq = shift_pos_array(x.seq, n)
    res.stat.hold_lengths = x.stat.hold_lengths[n:]
    res.stat.exit_delays = x.stat.exit_delays[n:]
    return res

def pos_pad(x, n=1):
    res = empty_pos()
    res.pos = pad_pos_array(x.pos, n)
    res.seq = pad_pos_array(x.seq, n)
    res.stat.hold_lengths = x.stat.hold_lengths
    res.stat.exit_delays = x.stat.exit_delays
    return res

def resample_returns(ret_df, months_per_period):
    if months_per_period >1:
        df = ret_df +1 # prep for geometric average calc
        df.columns = range(len(df.columns))
        df = df.assign(qtr = pd.array(pd.Series(range(len(df))) // months_per_period)) # assign quarter #
        # pandas will, unfortunately, sum NAs as though they are 0's, and multiply NAs as though they are 1's
        # We are forced to do our own NA/nan handling. Fortunately it's not too involved
        na_arr = np.tile(0, df.shape)
        na_arr[df.isna()] = 1
        na_df = pd.DataFrame(na_arr)
        na_df['qtr'] = df['qtr'].to_numpy()
        # df[:,:] assignment proved to be too slow
        #na_df = df.copy()
        #na_df.iloc[:,:] = 0
        #na_df[df.isna()] = 1
        na_idx = na_df.groupby(['qtr']).max() # will be =1 if there is at least one NA/nan

        r = df.groupby(['qtr']).prod()-1
        r[na_idx==1] = np.nan
        return r
    return ret_df


def flt_str_expand(val_map, separator=' '):
    shp = first_item_shape(val_map)
    res = pd.DataFrame(np.empty(shp, dtype=str))
    res_non_empty = np.tile(False, shp)
    for key, flt in val_map.items():
        sep_df = pd.DataFrame(res_non_empty & flt).replace({False:'', True:separator})
        val_df = pd.DataFrame(flt).replace({False:'', True:key})
        res = res + sep_df + val_df
        res_non_empty = res_non_empty | flt
    return res.to_numpy()

def get_model_spec(flt, cfg, wgt_list):
    res = []
    res.append('Entry conditions (AND):')
    res.append('  (Enter when entry signal = 1 and exit signal =0)')
    efk=list(flt.entry_map.keys())
    efk.sort
    for k in efk:
        res.append('* {}'.format(flt.entry_map[k].dscr))
    res.append('Allow entry cancellation if exit triggers: {}'.format(cfg.inv.entry.allow_cancel))

    res.append('')
    res.append('Exit conditions (OR):')
    res.append('  If entered less than {} months ago, do not exit'.format(cfg.inv.hold_months))

    efk=list(flt.exit_map.keys())
    efk.sort
    for k in efk:
        res.append('* {}'.format(flt.exit_map[k].dscr))

    res.append('')
    period_length_str = "month" if cfg.inv.months_per_period ==1 else "quarter" if cfg.inv.months_per_period ==3 else str(cfg.inv.months_per_period) + "mo"
    res.append('Period length: 1 {}'.format(period_length_str))

    res.append('')
    res.append('Portfolio holdings rule:')
    res.append('- Can enter in {} {}s'.format(cfg.inv.entry.delay, period_length_str))
    if cfg.inv.exit.is_uniform_delay:
        exit_delay_str = 'Can exit in {} {}s'.format(cfg.inv.exit.delay, period_length_str)
    else:
        exit_delay_str = 'Exit delay depends on fund (plus {} period(s) internal processing delay)'.format(cfg.inv.exit.processing_time)
    res.append('- {}'.format(exit_delay_str))

    res.append('')
    res.append('Additional filters (these don\'t change with time):')
    res.append('* Currency filter: {}'.format(cfg.ccy_flt))
    res.append('* Geographical mandate in {"North America", "Global", "Europe"}')
    res.append('* Include "dead" funds')

    res.append('')
    weighting_scheme_str = "Weighting scheme:"
    def fmt(v):
        return "_" if v is None else "{:.1f}".format(v*100)
    if cfg.inv.weigh_by_strategy:
        res.append(weighting_scheme_str)
        for w in cfg.inv.weight_cfg:
            res.append('* {}: {}-{}%'.format(w.inv_strats
                , fmt(w.min_weight), fmt(w.max_weight)))
        last_wgt_row = wgt_list[-1]
        res.append('* <Remaining funds>: {}-{}%'.format(
              fmt(last_wgt_row[0]), fmt(last_wgt_row[1])
        ))
    else:
        res.append(weighting_scheme_str + " Uniform")
    return res

def bool_mat_to_str_ind(v, str_ind):
    ret = np.tile('', v.shape)
    ret[v] = str_ind
    return ret

def validate_missing_returns(miss_flg, fund_ids, ever_entered_flt):
    if miss_flg.size <=0:
        return
    # This method will raise an error if the following condition fails for any of the funds:
    # If we have returns for some period(s), following by period(s) of missing returns,
    # make sure the reamining future returns are also missing.
    # That is, there are no "holes" in return series: only leading and trailing "blanks" with regard to time.
    v= np.tile(0, miss_flg.shape[1]) # STAGE 0: Allow as many missing returns as there are, until returns start coming
    for t in range(miss_flg.shape[0]):
        cur_missing = miss_flg[t,:].ravel()
        v[(cur_missing=='1') & ((v==0) | (v==2))] += 1 # Stage 1: returns started coming
        v[(cur_missing=='x') & (v==1)] += 1 # Stage 2: returns stopped coming
    if np.max(v) >=3:
        raise RuntimeError("Holes in returns for funds: {}, ...".format(
                fund_ids[ever_entered_flt][np.where(v==2)][:5]))

def add_ret_chart(ws_sum,  ws_ret, anchor_cell, df):
    ret_chart = opxc.ScatterChart()
    ret_chart.title = "Portfolio Cumulative Return (raw)"
    #ret_chart.style = 13
    ret_chart.x_axis.title = 'Period'
    ret_chart.y_axis.title = 'Return'

    xvalues = opxc.Reference(ws_ret, min_col=1, min_row=2, max_row=len(df)+1)
    yvalues = opxc.Reference(ws_ret, min_col=3, min_row=2, max_row=len(df)+1)
    series = opxc.Series(yvalues, xvalues, title='Return')
    ret_chart.series.append(series)

    ret_chart.legend = None
    ret_chart.width = 20
    ret_chart.height = 12
    ret_chart.x_axis.scaling.min = opxd.to_excel(df.index[0])-5
    ret_chart.x_axis.scaling.max = opxd.to_excel(df.index[len(df)-1])+5
    ret_chart.x_axis.majorUnit = 365.25*2
    ret_chart.x_axis.minorUnit = 365.25/2
    ret_chart.x_axis.minorGridlines = opxca.ChartLines()
    dashed_line = opxdl.LineProperties(prstDash='dash')
    ret_chart.x_axis.minorGridlines.graphicalProperties = opxcs.GraphicalProperties()
    ret_chart.x_axis.minorGridlines.graphicalProperties.line = dashed_line 
    ret_chart.x_axis.number_format ='mm/yyyy'

    ws_sum.add_chart(ret_chart, anchor_cell)

def prep_run_summary(portf_ret_df, stat, weights, rpt):
    vol_ann = np.std(portf_ret_df['Raw_Ret'], ddof=1)*np.sqrt(12)
    mean_ret_ann = np.average(portf_ret_df['Raw_Ret'])*12
    sr_ann = mean_ret_ann / vol_ann

    hold_lengths = list( itertools.chain.from_iterable(stat.hold_lengths) )
    exit_delays = list( itertools.chain.from_iterable(stat.exit_delays) )

    def add_row(r, k, v):
        r.append((k + ': ',v))
    rows = []
    add_row(rows, 'Sharpe ratio (full sample)', sr_ann)
    add_row(rows, 'Volatility (annualized)', vol_ann)
    add_row(rows, 'Mean return (annualized)', mean_ret_ann)
    add_row(rows, 'Avg number of holdings per period', np.average(np.sum(weights>0,1)))
    add_row(rows, 'Avg number of periods a fund is held', round(np.average(hold_lengths), 2))
    add_row(rows, 'Positions in last period', round(np.sum(weights[-1, :]>0)))
    add_row(rows, 'Funds ever entered', sum(sum(weights>0)>0))
    add_row(rows, 'Total redemptions', len(hold_lengths))
    add_row(rows, 'Funds with missing returns', len(rpt.miss_ret_rpt.df))
    add_row(rows, 'Avg exit delay (periods)', round(np.average(exit_delays), 2))

    return pd.DataFrame(rows, columns=['key', 'value'])

def write_run_summary_xl(ws_sum, df):
    ws_sum.column_dimensions["B"].width = 30
    ws_sum.column_dimensions["D"].width = 30

    num_fmts = ['#,##0.00','#,##0.0000', '#,##0.0000', '#,##0.0', '#,##0.0'
                , None, None, None, None, '#,##0.0']

    key_col = "B"
    val_col = "C"
    ln=1
    for rn, row in enumerate(df.itertuples(), 1):
        if rn==6:
            key_col = "D"
            val_col = "E"
            ln = 1
        ws_sum["{}{}".format(key_col, ln)] = row.key
        ws_sum["{}{}".format(val_col, ln)] = row.value
        fmt = num_fmts[rn-1]
        if fmt is not None:
            ws_sum["{}{}".format(val_col, ln)].number_format = fmt
        ln += 1

    for i in range(1,5):
        ws_sum["B{}".format(i)].alignment = opxsa.Alignment(horizontal='right')
        ws_sum["D{}".format(i)].alignment = opxsa.Alignment(horizontal='right')

def combine_filter_map(filters, op):
    if op == operator.or_:
        seed_val = False
    elif op == operator.and_:
        seed_val = True
    else:
        raise RuntimeError("Operator nto supported: {}".op)
    res = np.tile(seed_val, next(iter(filters.values())).filter.shape)
    for k, v in filters.items():
        res = op(res, v.filter)
    return res

def top_weight_rows(df, min_rows, max_rows, weight_fld, min_weight):
    df = df.sort_values(by=weight_fld, ascending=False)
    df = df.iloc[:max_rows,:]
    if len(df)>=min_rows:
        if df.iloc[min_rows-1][weight_fld] <= min_weight:
            return df.iloc[:min_rows,:]
        else:
            return df[df[weight_fld] >= min_weight]
    else:
        return df

_bold_style = opxs.Font(bold=True) 

def write_fld_sum(ws, df, fld_dscr, starting_row = 1):
    title_cell = openpyxl.cell.cell.Cell(ws, row=starting_row, column=1, value='Most common {}'.format(fld_dscr))
    title_cell.font = _bold_style
    ws._cells[(starting_row, 1)] = title_cell
    my_opx_append_rows(ws, dataframe_to_rows(df, index=False), number_formats=['0.0%'], hdr_font=_bold_style, starting_row = starting_row+1)

def output_std_rpt(wb, rpt, sheet_name):
    ws = wb.create_sheet(title=sheet_name)
    freeze_col = opxu.get_column_letter(len(rpt.ref_flds)+1) 
    ws.freeze_panes = '{}2'.format(freeze_col)
    my_opx_append_rows(ws, dataframe_to_rows(rpt.df, header=True, index=False), hdr_font=_bold_style)

def portf_returns(pos, ret, flt=None):
    if flt is not None:
        pos = pos[:,flt]
        ret = ret[:,flt]

    ret[pos==0] = 0 # mask nan returns for the periods when we have no positions
    return np.einsum('ij, ij->i', ret, pos)

def sum_val_by_period(weights, ref, fld, top_rows=None):
    ids = ref.fund_ids[ref.fund_flt]
    def sum_pos(pos, ref_col, fld, ids):
        df = pd.DataFrame(pos.transpose(), columns = ref.pos_dates, index = ids)
        prd_dates = pd.date_range(np.min(df.columns)+ pd.offsets.QuarterEnd(1), np.max(df.columns), freq='Q')
        # quarter ends only
        df = df.loc[:,prd_dates]
        c = df.columns.copy()
        df[fld] = ref_col
        dfu = pd.melt(df,id_vars = [fld], value_vars=c)
        dfu=dfu.rename(columns={'variable': 'dt'})
        dfu = dfu[dfu.value>0]
        cnt_df = dfu.groupby(['dt',fld]).sum().reset_index().sort_values(['dt','value', fld], ascending=False)
        cnt_df['rank'] = cnt_df.groupby(['dt']).rank('first', ascending=False)

        if top_rows is not None:
            cnt_df = cnt_df[cnt_df['rank'] <= top_rows]
        y=pd.pivot_table(cnt_df, values=['value', fld], index=['rank'], columns=['dt'], aggfunc=np.min)
        z = y.reorder_levels([1, 0], axis=1).sort_index(axis=1)
        z.columns = z.columns.set_levels(z.columns.levels[0].strftime('%Y-%m'), level=0)
        return z

    ref_col = ref.df.loc[ids, [fld]]
    res = argparse.Namespace()
    res.cnt_df = sum_pos(weights>0, ref_col, fld, ids)
    res.frac_df = sum_pos(weights, ref_col, fld, ids)
    return res

# Assumption: index are consecutive integers
def write_df_to_csv(df, fn, highlight_lines=None, csv_params=None):
    df = df.copy()
    if highlight_lines:
        highlight_lines = highlight_lines.copy()
        highlight_lines.sort(reverse=True)
    for l in highlight_lines:
        df.loc[l-0.5] = [np.nan] * len(df.columns)
    df = df.sort_index().reset_index(drop=True)
    df.to_csv(fn, **csv_params)

def calc_uniform_weights(pos):
    tot_funds = np.sum(pos,1)
    tot_funds[tot_funds==0] = np.nan # Prevent division by zero
    return np.einsum('ij, i->ij', pos, np.divide(1.0, tot_funds))    

# Important assumption: all rows of df have non-zero weight
# It is also expected that df.wgt adds up to 1.0 (this is the invariant preserved by the function)
def apply_weight_bounds(df):
    df = df.copy().assign(br = 0)
    df = df.assign(w_orig = df.wgt) # for debugging only; not used by the algorithm
    rows_to_process = np.tile(True, len(df))
    iter = 0
    while np.sum(rows_to_process)>0:
        iter +=1
        low_bound_breached = df.wgt < df.style_wgt_min
        high_bound_breached = df.wgt > df.style_wgt_max
        not_breached = ~ (low_bound_breached | high_bound_breached)
        rows_to_process = [False] # [np.tile(False, len(df))]
        if sum(~not_breached) >0:
            lb_df = df.loc[low_bound_breached]
            hb_df = df.loc[high_bound_breached]
            excess_wgt = np.sum(hb_df.wgt - hb_df.style_wgt_max) - np.sum(lb_df.style_wgt_min - lb_df.wgt)

            # enforce bounds
            df.loc[low_bound_breached, 'wgt'] = lb_df.style_wgt_min
            df.loc[high_bound_breached, 'wgt'] = hb_df.style_wgt_max
            df.loc[low_bound_breached, 'br'] = -1
            df.loc[high_bound_breached, 'br'] = 1

            # re-distribute excess weight
            if abs(excess_wgt) > SMALL_FLOAT:
                #assert np.sum(not_breached) >0, "Excess weight found with no rows to distribute it to"
                if np.sum(not_breached) <=0:
                    # There are no rows remaining that have no limits breached, but we still have excess weight
                    # This is possible, and is not necessarily an error.
                    # First determine what bounds we want to move "a little bit": upper or lower.
                    # This depends on the sign of the excess weight
                    fudge_rows = df.br == -np.sign(excess_wgt)
                    df.loc[fudge_rows, 'wgt'] += excess_wgt / np.sum(fudge_rows)

                    # See how we did on not breaching constraints during the last "weight redistriution" step
                    new_breach = (df.wgt < df.style_wgt_min) | (df.wgt > df.style_wgt_max)
                    if np.sum(new_breach)>0:
                        # Note that being here does NOT necessarily mean that the problem is incorrectly stated
                        # If the bounds are extremely tight, this algorithm may not be able to find a valid solution.
                        # This will likely happen on a "truncated" problem (that is, when not all styles have
                        # a position).
                        # In this case, it is likely a good idea to proceed with the solution found here rather
                        # than making method even more complicated. We raise an exception for now, because if this
                        # ever comes up, we want to look at the specific case before starting to ignore the error.
                        raise Exception("Failed to find satisfactory weights")
                else:
                    df.loc[not_breached, 'wgt'] = df.loc[not_breached, 'wgt'] + excess_wgt / np.sum(not_breached)
                    rows_to_process = not_breached
    assert abs(np.sum(df.wgt) - np.sum(df.w_orig)) < SMALL_FLOAT, "Algorithm invariant broken when redistributing weights"
    return df

def calc_weights(pos, grp_arr, style_wgt_range):
    ret = np.tile(np.nan, pos.shape)
    tot_funds = pos.shape[1]
    has_pos = pos>0
    style_weights_df = pd.DataFrame(style_wgt_range, index=range(len(style_wgt_range)), columns=['style_wgt_min', 'style_wgt_max'])
    pos_base_wgt_df = pd.DataFrame(grp_arr, columns=['style'])
    style_weights_sum = style_weights_df.sum()
    style_weights_sum_mid = np.average(style_weights_sum)
    rng_flds = ['style_wgt_min', 'style_wgt_max']
    for t in range(pos.shape[0]):
        styles = grp_arr[has_pos[t,:].ravel()]
        if len(styles) <=0:
            ret[t,:] = np.nan
        else:
            weights_df = pd.DataFrame({0:styles,'cnt':1}).groupby(0).count()
            assert len(weights_df)>0, "Unexpected error in style weighting"
            weights_df = weights_df.join(on=0, how='left', other=style_weights_df)
            if len(weights_df) < len(style_weights_df):
                # Not all styles present: scale up the config; i.e. "normalize"
                # weights_df_orig = weights_df.copy() # comment out for debugging
                
                observed_style_weights_sum = weights_df[rng_flds].sum()
                # Uniform scaling to preserve non-range configs where max=min
                weights_df[rng_flds] *= style_weights_sum_mid / np.average(observed_style_weights_sum)
                # older code: separate scaling factor for max/min points
                #weights_df[rng_flds] /= observed_style_weights_sum / style_weights_sum
            weights_df = apply_weight_bounds(weights_df.assign(wgt = weights_df.cnt / weights_df.cnt.sum()))
            weights_df.wgt /= weights_df.cnt

            #weights_df = weights_df.assign(wgt = weights_df.style_wgt / weights_df.cnt)
            pos_wgt_df = pos_base_wgt_df.copy()
            pos_wgt_df.loc[~has_pos[t,:], 'style'] = -100000
            pos_wgt_df = pos_wgt_df.join(on='style', other=weights_df[['wgt']])
            pos_wgt_df[pd.isna(pos_wgt_df.wgt)] = 0
            assert len(pos_wgt_df) == tot_funds, "Some records were lost when calculating weights"
            ret[t,:] = pos_wgt_df.wgt
    return ret

def build_pos(flt, inv_cfg, entry_delay, exit_delay):
    en_countdown = entry_delay+1
    ex_countdown = exit_delay+1
    en_flt = flt.entry & ~flt.exit # do not enter if both entry and exit signals are true
    pos_seq = np.tile(np.nan, en_flt.shape)
    cur_pos = np.tile(0, (1,en_flt.shape[1]))
    en_ctrl = cur_pos.copy()
    ex_ctrl = cur_pos.copy()
    res = empty_pos()
    for t in range(flt.entry.shape[0]):
        # Recordkeeping section
        # keep track of how many periods we were in each fund
        cur_pos[cur_pos!=0] +=  1
        # "decay" the wait time on all entry signals by 1
        en_ctrl[en_ctrl>0] -= 1
        # "decay" the wait time on all exit signals by 1
        ex_ctrl[ex_ctrl>0] -= 1

        # Process entry signal
        new_enter_trig = en_flt[t,:] & (en_ctrl==0) & (cur_pos<=0) # all funds with entry signal, which we did not start "entering" yet
        en_ctrl[new_enter_trig] = en_countdown[new_enter_trig]
        if inv_cfg.entry.allow_cancel:
            en_ctrl[:,flt.exit[t,:]] = 0
        # enter
        cur_pos[(cur_pos<=0) & (en_ctrl==1)] = 1

        # Process exit signal
        # Exit signal is true and we spent the specified # of periods in a position
        new_exit_trig = flt.exit[t,:] & (ex_ctrl==0) & (cur_pos > (inv_cfg.hold_months // inv_cfg.months_per_period) - exit_delay)
        ex_ctrl[new_exit_trig] = ex_countdown[new_exit_trig]
        exiting = (cur_pos>0) & (ex_ctrl==1)
        res.stat.hold_lengths.append(cur_pos[exiting])
        res.stat.exit_delays.append(exit_delay[exiting])
        cur_pos[exiting] = 0

        # a little extra safety: may happen if entry/exit delays are too unbalanced.
        en_ctrl[cur_pos>0] = 0
        ex_ctrl[cur_pos<=0] = 0

        pos_seq[t,:] = cur_pos
    res.seq = pos_seq
    res.pos = pos_seq.copy()
    res.pos[res.pos>0] = 1
    return res

def calc_weights_by_style(ref_sub_df, weight_cfg, pos, style_fld):
    ref_sub_df = ref_sub_df.assign(grp_num=len(weight_cfg)) # Map to remider items by default
    #wgt_predetermined = 0
    wgt_min = 1
    wgt_max = 1
    for cf in weight_cfg:
        #if cf.min_weight == cf.max_weight:
        #    wgt_predetermined -= cf.max_weight
        #else:
        if cf.min_weight is not None:
            wgt_max -= cf.min_weight
        if cf.max_weight is not None:
            wgt_min -= cf.max_weight
        else:
            wgt_min = 0 # just one unrestricted max means there is no required min for the leftover items
    assert wgt_max >0, "Min weight restirctions are too heavy (nothing is left for the remaining items)"
    wgt_min = np.max([0.0, wgt_min])
    wgt_min = None if abs(wgt_min) < SMALL_FLOAT else wgt_min
    wgt_max = None if abs(wgt_max-1) < SMALL_FLOAT else wgt_max

    style_wgt_range =[(cf.min_weight, cf.max_weight) for cf in weight_cfg]
    style_wgt_range.append((wgt_min, wgt_max)) # Restrictions for remainder items, if any
    for cf_idx in range(len(weight_cfg)):
        cf = weight_cfg[cf_idx]
        ref_sub_df.loc[ref_sub_df[style_fld].isin(cf.inv_strats), ['grp_num']] = cf_idx
    # Each item should have a non-empty list of funds corresponding to it, and there should be
    # some funds left to get the remaining weight
    assert len(ref_sub_df['grp_num'].unique()) == len(weight_cfg) +1, 'Style weighting config is inconsistent'
    style_grp_arr = ref_sub_df['grp_num'].to_numpy()
    weights = calc_weights(pos, style_grp_arr, style_wgt_range)
    if not np.all(np.abs(np.sum(weights,1)[np.sum(pos, 1)>0]-1.0)<SMALL_FLOAT):
        x=1
    assert np.all(np.abs(np.sum(weights,1)[np.sum(pos, 1)>0]-1.0)<SMALL_FLOAT), "Weights for some periods don't add up to 1"
    return weights, style_wgt_range

def enrich_ref(ref_df, ml):
    ref_df = ref_df.copy()
    ref_df.insert(loc=1, column='trackRecYrs', value=((ml.calc.dates[-1] - ref_df.INCEPTION_DATE).dt.days/365.25).ravel())
    ref_df.insert(loc=1, column='AUM', value=ml.db_data.aumTS.loc[ml.calc.dates[-1]].to_numpy().ravel()/(1000*1000))
    ref_df.insert(loc=1, column='paVol_1y', value=ml.calc.pAlphaVol_12mo[-1,:].ravel())
    ref_df.insert(loc=1, column='paSR_1y', value=ml.calc.pAlphaSrp_12mo[-1,:].ravel())
    ref_df.insert(loc=1, column='Beta', value=ml.calc.betas[-1,:].ravel())
    return ref_df

def enrich_df(df, ref_df, dtp, val_map=None, flt=None,  add_flds=None, flds=None):
    if flt is not None: df = df.loc[flt,:] 
    if val_map is None:
        if dtp ==bool: 
            val_map = {False:np.nan, True:1}
    if val_map is not None:
        df = df.replace(val_map)

    if flds is None:
        ref_flds = ['FUND_NAME', 'Beta', 'paSR_1y', 'paVol_1y', 'AUM', 'trackRecYrs']
    else:
        ref_flds = ['FUND_NAME'] + flds
    if add_flds is not None:
        ref_flds.extend(add_flds)
    ref_subset_df = ref_df[ref_flds]

    out_df = pd.merge(ref_subset_df,df,left_index=True, right_index=True)
    assert (len(df) == len(out_df)), 'Lost some records during enrichment transformation(s)!'
    ret = argparse.Namespace()
    ret.df = out_df.reset_index()
    ret.ref_flds = ref_flds.copy()
    return ret

def save_model_spec_code(main_fn, output_path):
    code_fn = "{}/util/sig.py".format(os.path.dirname(main_fn))
    with open(code_fn) as f:
        model_spec_code = [line.rstrip() for line in f]

    model_spec_code = model_spec_code[
        model_spec_code.index('# Model Spec: Begin'): model_spec_code.index('# Model Spec: End')+1]
    np.savetxt("{}/model_spec.py".format(output_path), model_spec_code, '%s')

def build_fld_summary(df, fld):
    df = df[[fld, 'wgt']].groupby([fld]).sum()
    df = top_weight_rows(df, 5, 10, 'wgt', 0.05).reset_index()
    return df[[df.columns[1], df.columns[0]]]

def enrich_and_write_ex(arr, ref, output_path, file_name, columns=None, **kwargs):
    columns = coalesce(columns, ref.dates.strftime('%Y-%m'))
    df = pd.DataFrame(arr.transpose()
        , columns = columns, index = ref.fund_ids[ref.fund_flt])
    ret = enrich_df(df, ref.df, arr.dtype, **kwargs)
    ret.df.to_csv("{}/{}".format(output_path, file_name), index=False, date_format='%Y-%m-%d')
    return ret

